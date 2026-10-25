"""Train, evaluate, and present the Week 4 heart disease capstone."""

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    ConfusionMatrixDisplay,
    accuracy_score,
    f1_score,
    precision_score,
    recall_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data" / "heart_disease_uci.csv"
OUTPUT = ROOT / "outputs"
SEED = 42
LABELS = ["No disease", "Disease"]


def make_preprocessor(X_train, scale_numeric):
    numeric = X_train.select_dtypes(include="number").columns
    categorical = X_train.select_dtypes(exclude="number").columns
    numeric_steps = [("imputer", SimpleImputer(strategy="median", add_indicator=True))]
    if scale_numeric:
        numeric_steps.append(("scaler", StandardScaler()))
    return ColumnTransformer(
        [
            ("numeric", Pipeline(numeric_steps), numeric),
            (
                "categorical",
                Pipeline(
                    [
                        ("imputer", SimpleImputer(strategy="most_frequent")),
                        ("encoder", OneHotEncoder(handle_unknown="ignore")),
                    ]
                ),
                categorical,
            ),
        ]
    )


def build_presentation(metrics, dataset_info):
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Heart Disease Prediction"
    slide.placeholders[1].text = (
        "Week 4 Machine Learning Capstone\n"
        "Logistic Regression and Decision Tree · UCI Heart Disease"
    )

    def bullets(title, lines):
        page = deck.slides.add_slide(deck.slide_layouts[1])
        page.shapes.title.text = title
        frame = page.placeholders[1].text_frame
        frame.clear()
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(22)
            paragraph.space_after = Pt(12)
        return page

    missing = dataset_info["missing_values"]
    top_missing = sorted(missing.items(), key=lambda pair: pair[1], reverse=True)[:3]
    bullets(
        "Dataset and target",
        [
            f"{dataset_info['rows']} records and {dataset_info['columns_count']} original columns.",
            "Target num was converted to binary: 0 = no disease; 1–4 = disease.",
            f"Class counts: {dataset_info['class_distribution']}.",
            "Dropped patient ID and collection site to avoid identifier and site shortcuts.",
            "Most missing: " + ", ".join(f"{name} ({count})" for name, count in top_missing) + ".",
            "Source: UCI Heart Disease, CC BY 4.0; see the README for attribution.",
        ],
    )
    bullets(
        "Preprocessing and evaluation",
        [
            "One reproducible 80/20 stratified train/test split (random seed 42).",
            "Numeric missing values: training-set median imputation plus missing indicators.",
            "Categorical values: training-set most-frequent imputation and one-hot encoding.",
            "Numeric scaling is fitted for Logistic Regression; Decision Trees use unscaled values.",
            "All preprocessing is inside each model pipeline, fitted on training rows only.",
            "Metrics: accuracy, precision, recall, F1, ROC-AUC, and confusion matrices.",
        ],
    )

    page = deck.slides.add_slide(deck.slide_layouts[5])
    page.shapes.title.text = "Test-set results"
    table = page.shapes.add_table(3, 6, Inches(0.45), Inches(1.7), Inches(12.4), Inches(1.7)).table
    headers = ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"]
    for col, label in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = label
        for run in cell.text_frame.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(14)
    for row, item in enumerate(metrics, start=1):
        values = [item["model"], *(f"{item[key]:.3f}" for key in ("accuracy", "precision", "recall", "f1", "roc_auc"))]
        for col, value in enumerate(values):
            table.cell(row, col).text = value
            for paragraph in table.cell(row, col).text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(14)

    page = deck.slides.add_slide(deck.slide_layouts[5])
    page.shapes.title.text = "Confusion matrices and ROC curves"
    page.shapes.add_picture(str(OUTPUT / "confusion_matrices.png"), Inches(0.4), Inches(1.25), width=Inches(6.2))
    page.shapes.add_picture(str(OUTPUT / "roc_curves.png"), Inches(6.8), Inches(1.25), width=Inches(6.1))

    page = deck.slides.add_slide(deck.slide_layouts[5])
    page.shapes.title.text = "Model comparison"
    page.shapes.add_picture(str(OUTPUT / "model_comparison.png"), Inches(1.0), Inches(1.3), width=Inches(11.3))

    winner = max(metrics, key=lambda item: (item["roc_auc"], item["f1"]))
    bullets(
        "Conclusion and limitations",
        [
            f"On this split, {winner['model']} has the strongest ROC-AUC ({winner['roc_auc']:.3f}) among these two baselines.",
            "The combined cohort spans multiple hospitals and has substantial missingness; results may reflect cohort differences.",
            "A single random split is educational evidence, not external clinical validation.",
            "This model is not a clinical diagnostic tool and must not guide patient care.",
        ],
    )
    deck.save(OUTPUT / "heart_disease_capstone.pptx")


def main():
    if not DATA.exists():
        raise SystemExit(f"Dataset not found: {DATA}")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    data = pd.read_csv(DATA, na_values=["?"])
    target = (pd.to_numeric(data["num"]) > 0).astype(int)
    # Exclude row IDs and hospital/source labels, which can create shortcuts.
    X = data.drop(columns=["id", "dataset", "num"])
    X_train, X_test, y_train, y_test = train_test_split(
        X, target, test_size=0.2, random_state=SEED, stratify=target
    )

    models = {
        "Logistic Regression": Pipeline(
            [
                ("preprocess", make_preprocessor(X_train, scale_numeric=True)),
                ("model", LogisticRegression(max_iter=2000, random_state=SEED)),
            ]
        ),
        "Decision Tree": Pipeline(
            [
                ("preprocess", make_preprocessor(X_train, scale_numeric=False)),
                ("model", DecisionTreeClassifier(random_state=SEED)),
            ]
        ),
    }

    results, predictions = [], {}
    for name, pipeline in models.items():
        pipeline.fit(X_train, y_train)
        predicted = pipeline.predict(X_test)
        probabilities = pipeline.predict_proba(X_test)[:, 1]
        predictions[name] = (predicted, probabilities)
        results.append(
            {
                "model": name,
                "accuracy": accuracy_score(y_test, predicted),
                "precision": precision_score(y_test, predicted, zero_division=0),
                "recall": recall_score(y_test, predicted, zero_division=0),
                "f1": f1_score(y_test, predicted, zero_division=0),
                "roc_auc": roc_auc_score(y_test, probabilities),
            }
        )

    metrics = pd.DataFrame(results)
    metrics.to_csv(OUTPUT / "model_metrics.csv", index=False)
    pd.DataFrame(
        {
            "actual": y_test.to_numpy(),
            **{f"{name}_predicted": pair[0] for name, pair in predictions.items()},
            **{f"{name}_probability": pair[1] for name, pair in predictions.items()},
        }
    ).to_csv(OUTPUT / "test_predictions.csv", index=False)

    info = {
        "rows": len(data),
        "columns_count": len(data.columns),
        "columns": data.columns.tolist(),
        "dtypes": {name: str(dtype) for name, dtype in data.dtypes.items()},
        "missing_values": {name: int(count) for name, count in data.isna().sum().items()},
        "raw_num_distribution": {str(k): int(v) for k, v in data["num"].value_counts().sort_index().items()},
        "class_distribution": {"no disease (0)": int((target == 0).sum()), "disease (1)": int((target == 1).sum())},
        "features_used": X.columns.tolist(),
        "categorical_features": X.select_dtypes(exclude="number").columns.tolist(),
        "numeric_features": X.select_dtypes(include="number").columns.tolist(),
    }
    (OUTPUT / "dataset_summary.json").write_text(json.dumps(info, indent=2), encoding="utf-8")

    fig, axes = plt.subplots(1, 2, figsize=(10, 4.5))
    for axis, (name, (predicted, _)) in zip(axes, predictions.items()):
        ConfusionMatrixDisplay.from_predictions(
            y_test, predicted, labels=[0, 1], display_labels=LABELS,
            cmap="Blues", values_format="d", ax=axis, colorbar=False,
        )
        axis.set_title(name)
    fig.tight_layout()
    fig.savefig(OUTPUT / "confusion_matrices.png", dpi=160)
    plt.close(fig)

    fig, axis = plt.subplots(figsize=(7, 5))
    for name, (_, probabilities) in predictions.items():
        false_positive, true_positive, _ = roc_curve(y_test, probabilities)
        auc = roc_auc_score(y_test, probabilities)
        axis.plot(false_positive, true_positive, label=f"{name} (AUC={auc:.3f})")
    axis.plot([0, 1], [0, 1], "k--", label="Chance")
    axis.set(xlabel="False positive rate", ylabel="True positive rate", title="ROC curves")
    axis.legend(loc="lower right")
    axis.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(OUTPUT / "roc_curves.png", dpi=160)
    plt.close(fig)

    metrics.set_index("model")[["accuracy", "precision", "recall", "f1", "roc_auc"]].plot(
        kind="bar", ylim=(0, 1), figsize=(9, 5), ylabel="Score", title="Test-set model comparison"
    )
    plt.xticks(rotation=0)
    plt.grid(axis="y", alpha=0.25)
    plt.tight_layout()
    plt.savefig(OUTPUT / "model_comparison.png", dpi=160)
    plt.close()

    print(metrics.to_string(index=False, float_format=lambda value: f"{value:.4f}"))
    print(f"Dataset summary and predictions saved in {OUTPUT}")
    build_presentation(results, info)
    print(f"Editable presentation saved to {OUTPUT / 'heart_disease_capstone.pptx'}")


if __name__ == "__main__":
    main()
